import re
from docx import Document

def check_references(doc: Document) -> list[dict]:
    errors = []

    heading_index = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().lower() == "список использованных источников":
            heading_index = i
            break

    if heading_index is None:
        return [{"paragraph_index": 0, "error": "Список использованных источников не найден"}]

    list_nums = set()
    for p in doc.paragraphs[heading_index + 1:]:
        text = p.text.strip()
        if not text:
            continue

        matches = re.findall(r"\[(\d+)\]", text)
        if matches:
            for m in matches:
                list_nums.add(int(m))
            continue

        m = re.match(r"^(\d+)[\.\)]\s*", text)
        if m:
            list_nums.add(int(m.group(1)))

    if not list_nums:
        return [{"paragraph_index": heading_index, "error": "Не удалось определить номера источников в списке"}]

    used_order = []
    used_set = set()
    for p in doc.paragraphs[:heading_index]:
        for bracket in re.findall(r"\[([^\]]+)\]", p.text):
            cleaned = bracket.replace(" ", "").replace("–", "-")
            for part in cleaned.split(","):
                if "-" in part:
                    try:
                        start, end = map(int, part.split("-"))
                        for num in range(start, end + 1):
                            if num not in used_set:
                                used_set.add(num)
                                used_order.append(num)
                    except ValueError:
                        continue
                else:
                    try:
                        num = int(part)
                        if num not in used_set:
                            used_set.add(num)
                            used_order.append(num)
                    except ValueError:
                        continue

    for num in sorted(used_set):
        if num not in list_nums:
            errors.append({
                "paragraph_index": heading_index,
                "error": f"Источник [{num}] упомянут в тексте, но не указан в списке"
            })

    for num in sorted(list_nums):
        if num not in used_set:
            errors.append({
                "paragraph_index": heading_index,
                "error": f"Источник [{num}] указан в списке, но не использован в тексте"
            })

    for i in range(len(used_order) - 1):
        if used_order[i] > used_order[i + 1]:
            errors.append({
                "paragraph_index": heading_index,
                "error": f"Нарушен порядок первых упоминаний: [{used_order[i+1]}] после [{used_order[i]}]"
            })
            break

    return errors
